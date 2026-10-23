"""06_User_Flows.docx: each flow as numbered steps plus a simple box-and-arrow diagram."""

from pathlib import Path

from docx.shared import Cm
from PIL import Image, ImageDraw, ImageFont

from docstyle import contents, new_document, para, steps

DIAGRAMS = Path(__file__).parents[1] / "diagrams"
FONT = "C:/Windows/Fonts/segoeui.ttf"
BOLD = "C:/Windows/Fonts/segoeuib.ttf"

# (title, [(short box label, full step)])
FLOWS = [
    (
        "Guest start",
        [
            ("Open the site", "Open the site; the sign-in page shows."),
            ("Continue as guest", "Click **Continue as guest**."),
            ("Start page", "The start page opens with an empty first board."),
            ("Use the app", "Use every feature except sharing; a bar reminds that guest boards are removed after 7 days."),
            ("Sign up to keep", "Click **Sign up to save your boards** to keep them: the boards stay with the new account."),
        ],
    ),
    (
        "Sign up and verify",
        [
            ("Create an account", "On the sign-in page click **Create an account**."),
            ("Email + password", "Type the email and the password twice; Weak passwords are refused."),
            ("Check your inbox", "Click **Create account**; the Check your inbox page opens."),
            ("Look in Spam/Junk", "Open the confirmation email, usually in **Spam** or **Junk**; mark it Not spam."),
            ("Click the link", "Click the link (valid 24 hours); the account is confirmed and signed in."),
            ("Start page", "The start page opens."),
        ],
    ),
    (
        "Google sign-in",
        [
            ("Continue with Google", "Click **Continue with Google**."),
            ("Pick account", "Choose a Google account on Google's page."),
            ("Back to the app", "Google returns to the app; a new account is created or the existing one is linked."),
            ("Start page", "The start page opens; no confirmation email is needed."),
        ],
    ),
    (
        "Forgot password",
        [
            ("Forgot password?", "On the sign-in page click **Forgot password?**."),
            ("Type email", "Type the email and click **Send reset link**."),
            ("Open the email", "Open the email (check Spam or Junk) and click the link."),
            ("New password", "Type the new password twice and click **Save new password**."),
            ("Signed in", "Other devices are signed out; you are signed in."),
        ],
    ),
    (
        "Create a board",
        [
            ("New board", "Click **New board** on the top bar or the start page."),
            ("Pick a start", "Choose Blank, Sprint, Bug tracking, Personal, your own template, or Describe it to AI."),
            ("Name it", "Type the board name and pick a workspace (optional)."),
            ("Create board", "Click **Create board** (or **Create with AI**)."),
            ("Board opens", "The new board opens."),
        ],
    ),
    (
        "Cards and drag and drop",
        [
            ("Add card", "Click **Add card** under a column, type the title, press Enter."),
            ("Drag card", "Drag the card to another column (or Space, arrows, Space)."),
            ("Saved", "The move shows at once and is saved."),
            ("Complete", "Click the circle to mark it complete; the progress bar moves."),
            ("Undo delete", "Delete a card and click **Undo** within 5 seconds to bring it back."),
        ],
    ),
    (
        "Invite and accept",
        [
            ("Share", "The admin opens **Share** on the board."),
            ("Invite by email", "Types the email and clicks **Invite**; an email is sent."),
            ("Or send link", "Optionally **Copy link** or **Share on WhatsApp**."),
            ("Open the link", "The invited person opens the link and signs in with that email."),
            ("Accept", "They click **Accept invitation**; the board opens for them as Viewer."),
        ],
    ),
    (
        "Permissions and request access",
        [
            ("Viewer", "A new member is a Viewer and can only look."),
            ("Try to edit", "When they try to change something, a message asks them to request access."),
            ("Request access", "They click **Request access**; the admin gets an email."),
            ("Admin decides", "The admin chooses **Allow: Can update tasks**, **Allow: Full access** or **Decline**."),
            ("New permission", "The member's permission changes at once."),
        ],
    ),
    (
        "Public board and join",
        [
            ("Turn on link", "The admin ticks **Anyone with the link can view**."),
            ("Share link", "Copies the public board link and shares it."),
            ("View only", "Anyone opens it without signing in and sees the board read-only."),
            ("Join board", "A signed-in visitor clicks **Join board** and becomes a Viewer."),
            ("Turn off", "Unticking the box stops the link."),
        ],
    ),
    (
        "Assign cards",
        [
            ("Open card", "Open a card on a shared board."),
            ("Tick member", "Tick a member under **Assigned to**."),
            ("Notified", "The member gets a notification and an email."),
            ("Assigned to me", "They filter the board with **Assigned to me** to see their cards."),
        ],
    ),
    (
        "Meeting link",
        [
            ("Paste link", "In Share, paste a Zoom, Google Meet or Teams link."),
            ("Save link", "Click **Save link**; **Start meeting** appears on the board."),
            ("Send to members", "Click **Send the link to members**: each gets an email and a notification."),
            ("Join", "Members click **Start meeting** to join."),
        ],
    ),
    (
        "AI chat actions",
        [
            ("Ask AI", "Click **Ask AI** on a board."),
            ("Ask", "Ask a question or for a change, e.g. move a card to Done."),
            ("Proposal", "The assistant answers and proposes the change."),
            ("Confirm", "Click **Confirm** to apply it, or **Dismiss**."),
            ("Board updates", "The board changes only after Confirm."),
        ],
    ),
    (
        "Voice input",
        [
            ("Microphone", "Click the microphone in a text box, or **Voice command**."),
            ("Allow", "The first time, allow the microphone in the browser."),
            ("Speak", "Speak in English or Arabic; the words appear as you talk."),
            ("Check", "Check the text, or the proposed command."),
            ("Save or Confirm", "Save the text, or click **Confirm** to run the command."),
        ],
    ),
    (
        "Files",
        [
            ("Open card", "Open a card."),
            ("Choose files", "Drop files or click **Choose files** (up to 4 MB each, 10 per card)."),
            ("Uploaded", "The file shows in Attachments; click it to download."),
            ("Make cover", "For an image, **Make cover** shows it on the card."),
            ("Delete", "The x button deletes the file."),
        ],
    ),
    (
        "Focus timer",
        [
            ("Focus timer", "Click **Focus timer** on a board."),
            ("Pick length", "Choose 25/5 or 50/10 and a card to focus on."),
            ("Start focus", "Click **Start focus**; the time counts down at the top."),
            ("Break", "When focus ends, a sound plays and the break starts."),
            ("Saved", "The focus session is saved on the card."),
        ],
    ),
    (
        "Language switch",
        [
            ("Click Arabic", "Click **العربية** on the top bar."),
            ("Right to left", "The whole app turns Arabic and right-to-left."),
            ("Remembered", "The choice is saved with the account."),
            ("Emails and AI", "Emails and AI answers follow the language."),
            ("English", "Click **English** to switch back."),
        ],
    ),
    (
        "Delete account",
        [
            ("Account menu", "Click your initial at the top right."),
            ("Delete account", "Choose **Delete account**."),
            ("Confirm", "Confirm in the window that asks again."),
            ("Removed", "The account, its boards and files are removed for good."),
        ],
    ),
]


def diagram(title: str, labels: list[str]) -> Path:
    """Boxes left to right with arrows, wrapping onto a second row when needed."""
    font, bold = ImageFont.truetype(FONT, 26), ImageFont.truetype(BOLD, 26)
    box_w, box_h, gap, per_row = 300, 96, 70, 4
    rows = (len(labels) + per_row - 1) // per_row
    width = per_row * box_w + (per_row - 1) * gap + 40
    height = rows * box_h + (rows - 1) * 60 + 40
    img = Image.new("RGB", (width, height), "white")
    d = ImageDraw.Draw(img)
    centers = []
    for n, label in enumerate(labels):
        row, col = divmod(n, per_row)
        x, y = 20 + col * (box_w + gap), 20 + row * (box_h + 60)
        first, last = n == 0, n == len(labels) - 1
        fill = "#753991" if last else "#032147" if first else "#eef4fb"
        d.rounded_rectangle([x, y, x + box_w, y + box_h], 18, fill=fill, outline="#209dd7", width=3)
        f = bold if first or last else font
        tw = d.textlength(label, font=f)
        d.text((x + (box_w - tw) / 2, y + box_h / 2 - 17), label, font=f, fill="white" if first or last else "#032147")
        centers.append((x, y))
    for (x1, y1), (x2, y2) in zip(centers, centers[1:], strict=False):
        if y1 == y2:
            a, b, y = x1 + box_w + 6, x2 - 6, y1 + box_h // 2
            d.line([a, y, b, y], fill="#ecad0a", width=6)
            d.polygon([(b, y), (b - 16, y - 11), (b - 16, y + 11)], fill="#ecad0a")
        else:
            sx, sy, ex, ey = x1 + box_w // 2, y1 + box_h + 4, x2 + box_w // 2, y2 - 4
            mid = sy + 26
            d.line([sx, sy, sx, mid, ex, mid, ex, ey], fill="#ecad0a", width=6, joint="curve")
            d.polygon([(ex, ey), (ex - 11, ey - 16), (ex + 11, ey - 16)], fill="#ecad0a")
    DIAGRAMS.mkdir(parents=True, exist_ok=True)
    path = DIAGRAMS / f"{title.lower().replace(' ', '-')}.png"
    img.save(path)
    return path


def build(out: Path) -> None:
    doc = new_document("Task Manager", "User flows")
    para(doc, "How people move through the app, step by step. Each flow starts at the dark box and ends at the purple one.")
    contents(doc)
    groups = [("Accounts", FLOWS[0:4]), ("Boards and cards", FLOWS[4:6]), ("Working together", FLOWS[6:11]), ("Helpers", FLOWS[11:15]), ("Settings", FLOWS[15:])]
    n = 0
    for group, flows in groups:
        doc.add_heading(group, level=1)
        for title, items in flows:
            n += 1
            doc.add_heading(f"{n}. {title}", level=2)
            p = doc.add_paragraph()
            p.add_run().add_picture(str(diagram(title, [label for label, _ in items])), width=Cm(16.5))
            p.paragraph_format.keep_with_next = True
            steps(doc, [step for _, step in items])
    doc.save(str(out))
