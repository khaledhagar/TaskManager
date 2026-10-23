"""Write the project's Markdown (headings, lists, checkboxes, tables, code) into a Word document."""

import re

from docx.shared import Cm, Pt, RGBColor

from docstyle import GRAY, add_runs, simple_table

DONE = RGBColor(0x1B, 0x7F, 0x3B)
LINK = re.compile(r"\[([^\]]+)\]\([^)]+\)")


def plain(text: str) -> str:
    """Markdown links become their text."""
    return LINK.sub(r"\1", text)


def render(doc, markdown: str, *, top: int = 2) -> None:
    """Render Markdown; `top` is the Markdown heading level that becomes Heading 1."""
    lines = markdown.splitlines()
    i = 0
    while i < len(lines):
        line = lines[i]
        if line.startswith("```"):
            i = _code(doc, lines, i + 1)
            continue
        if line.lstrip().startswith("|"):
            i = _table(doc, lines, i)
            continue
        heading = re.match(r"^(#+)\s+(.*)", line)
        if heading:
            level = len(heading.group(1)) - top + 1
            if level >= 1:
                doc.add_heading(plain(heading.group(2)), level=min(level, 3))
        elif item := re.match(r"^(\s*)[-*]\s+(\[[ xX]\]\s+)?(.*)", line):
            _item(doc, len(item.group(1)) // 2, item.group(2), plain(item.group(3)))
        elif num := re.match(r"^\s*(\d+)\.\s+(.*)", line):
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Cm(0.8)
            p.paragraph_format.first_line_indent = Cm(-0.6)
            p.add_run(f"{num.group(1)}.  ").bold = True
            add_runs(p, plain(num.group(2)))
        elif line.strip():
            add_runs(doc.add_paragraph(), plain(line.strip()))
        i += 1


def _item(doc, depth: int, box: str | None, text: str) -> None:
    p = doc.add_paragraph(style="List Bullet 2" if depth else "List Bullet")
    p.paragraph_format.space_after = Pt(2)
    if box:
        done = box[1] in "xX"
        run = p.add_run("Done  " if done else "Open  ")
        run.bold = True
        run.font.size = Pt(8)
        run.font.color.rgb = DONE if done else GRAY
    add_runs(p, text)


def _code(doc, lines: list[str], i: int) -> int:
    while i < len(lines) and not lines[i].startswith("```"):
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.left_indent = Cm(0.5)
        run = p.add_run(lines[i])
        run.font.name = "Consolas"
        run.font.size = Pt(9)
        i += 1
    doc.add_paragraph()
    return i + 1


def _table(doc, lines: list[str], i: int) -> int:
    rows = []
    while i < len(lines) and lines[i].lstrip().startswith("|"):
        cells = [plain(c.strip()) for c in lines[i].strip().strip("|").split("|")]
        if not all(re.fullmatch(r":?-+:?", c) for c in cells):
            rows.append(cells)
        i += 1
    header, body = rows[0], rows[1:]
    width = 17 / len(header)
    widths = [width] * len(header)
    if len(header) > 1:
        lengths = [max(len(r[c]) if c < len(r) else 0 for r in rows) for c in range(len(header))]
        total = sum(min(n, 60) + 4 for n in lengths)
        widths = [17 * (min(n, 60) + 4) / total for n in lengths]
    body = [(r + [""] * len(header))[: len(header)] for r in body]
    simple_table(doc, header, body, widths)
    return i
