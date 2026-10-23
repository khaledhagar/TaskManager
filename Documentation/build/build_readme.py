"""Build Documentation/README.docx: the user guide (T-628). Run `uv run python build_readme.py`."""

from pathlib import Path

from docx.shared import Cm

from docstyle import bullets, contents, image, new_document, note_box, para, simple_table, steps

SITE = "https://task-manager-opal-xi-63.vercel.app"
OUT = Path(__file__).parents[1] / "README.docx"


def build() -> Path:
    doc = new_document("Task Manager", "User guide")
    note_box(
        doc,
        "Important: the confirmation email usually lands in Spam or Junk",
        [
            "When you create an account with your email address, Task Manager sends you a confirmation email. "
            "Most email services put it in the **Spam** or **Junk** folder, not in your inbox.",
            "Open your **Spam** or **Junk** folder, find the email from **Task Manager**, mark it as **Not spam**, "
            "then click the link inside. The link works for 24 hours.",
            "No email after a few minutes? On the **Check your inbox** page, click **Resend email**.",
        ],
    )
    para(
        doc,
        f"Task Manager is a simple board app for planning work. Open it at **{SITE}** in any modern browser "
        "(Chrome, Edge, Firefox or Safari) on a computer, tablet or phone. It is free to use.",
    )
    contents(doc)
    getting_started(doc)
    start_page(doc)
    create_board(doc)
    work_on_board(doc)
    inside_card(doc)
    ai_helpers(doc)
    find_cards(doc)
    voice(doc)
    focus_timer(doc)
    board_options(doc)
    share(doc)
    account(doc)
    look(doc)
    limits(doc)
    help_section(doc)
    doc.save(str(OUT))
    return OUT


def getting_started(doc):
    doc.add_heading("1. Getting started", level=1)
    doc.add_heading("Create an account with your email", level=2)
    steps(
        doc,
        [
            "Open the site and click **Create an account** under the sign-in form.",
            "Type your email and a password twice. The bar under the password shows its strength: **Weak** passwords "
            "are refused; use at least 10 characters with letters and numbers.",
            "Click **Create account**. The **Check your inbox** page opens.",
            "Open the confirmation email (look in **Spam** or **Junk**, see the note at the top) and click its link. You are signed in and your start page opens.",
        ],
    )
    image(doc, "02-create-account", "Create your account, with the password strength bar")
    image(doc, "03-check-your-inbox", "After Create account: check your inbox, and your Spam or Junk folder")
    doc.add_heading("Other ways in", level=2)
    bullets(
        doc,
        [
            "**Continue with Google**: signs you in with your Google account. No password and no confirmation email needed.",
            "**Continue as guest**: try the app without an account. Guest boards are removed after 7 days. "
            "To keep them, click **Sign up to save your boards** in the bar at the top.",
            "**Forgot password?** on the sign-in page emails you a link to choose a new password (this email may also land in Spam or Junk).",
        ],
    )
    image(doc, "01-sign-in", "The sign-in page: Google, email and password, or guest")
    image(doc, "37-guest", "A guest account shows a reminder bar at the top")


def start_page(doc):
    doc.add_heading("2. Your start page", level=1)
    para(doc, "Every time you sign in, the start page **Your work** opens. Click the app name at the top left to come back to it.")
    bullets(
        doc,
        [
            "**Left panel**: your workspaces. Under each one, **Boards** lists its boards and **Members** lists the people in them.",
            "**Middle**: boards grouped by workspace, then **Boards without a workspace**, **Shared with you** and **Starred**.",
            "**What's new**: invitations and changes to cards assigned to you.",
            "**New workspace** and **New board** buttons start something new.",
        ],
    )
    image(doc, "32-start-page", "The start page with a workspace and its boards")
    doc.add_heading("Workspaces", level=2)
    para(doc, 'A workspace groups related boards, for example "Website relaunch" with boards for Design and Development.')
    steps(
        doc,
        [
            "Click **New workspace** and type its name.",
            "Name each board. Each one starts with To Do, In Progress and Done; add, remove or rename columns, "
            "or click **Suggest columns with AI**. Click **Add another board** for more.",
            "Click **Create workspace**.",
        ],
    )
    image(doc, "05-new-workspace", "A new workspace with two boards")


def create_board(doc):
    doc.add_heading("3. Create a board", level=1)
    steps(
        doc,
        [
            "Click **New board** (top bar or start page).",
            "Choose how to start: **Blank**, **Sprint**, **Bug tracking**, **Personal**, one of your own templates, or **Describe it to AI**.",
            "Type the board name, pick a workspace if you like, and click **Create board**.",
        ],
    )
    image(doc, "06-new-board-templates", "Start from a template")
    doc.add_heading("Let AI build the board", level=2)
    para(
        doc,
        "Choose **Describe it to AI**, write one sentence about your project, and click **Create with AI**. "
        "In up to a minute you get a board with columns and starter cards.",
    )
    image(doc, "07-new-board-ai", "Describe the project in one sentence")
    image(doc, "08-board-made-by-ai", "The board the AI made for an online bakery")


def work_on_board(doc):
    doc.add_heading("4. Work on a board", level=1)
    image(doc, "09-board", "A board: columns, cards, labels, priorities, due dates and the progress bar")
    bullets(
        doc,
        [
            "**Add a card**: click **Add card** under a column, type the title, press Enter. A board holds up to 100 cards.",
            "**Move a card**: drag it to another place or column. With the keyboard: focus the card, press Space, use the arrow keys, press Space again.",
            "**Complete a card**: click the circle before its title. The progress bar at the top shows how many are done.",
            "**Columns**: click a column name to rename it; drag the dots on its left to reorder; **Add column** adds one. "
            "The **...** menu of a column can set a **WIP limit** (most cards allowed), move it left or right, archive or delete it.",
            "**Undo**: after deleting a card or column, click **Undo** in the message at the bottom.",
            "**Colors**: cards due soon or overdue are highlighted; the priority badge shows Low, Medium, High or Urgent.",
            "**Rename or star the board**: click its title to rename it; click the star to keep it under **Starred**.",
        ],
    )


def inside_card(doc):
    doc.add_heading("5. Inside a card", level=1)
    para(doc, "Click a card to open it. Everything is saved as you go.")
    bullets(
        doc,
        [
            "**Description**: write details. Markdown works: **bold**, lists, headings and links; click **Preview** to see it.",
            "**Checklist**: type an item and press Enter; tick items as you finish them.",
            "**Priority**, **Due date** and **Labels** are on the right. Create new labels with a name and a color.",
            "**Assigned to**: tick the board members who work on the card.",
            "**Attachments**: drop files or click **Choose files** (images, PDF, text, CSV or Office files, up to 4 MB each). **Make cover** shows an image on the card.",
            "**Comments**: write a comment; type **@** to mention a member, who gets a notification.",
            "**History** lists every change to the card.",
            "**Mark complete**, **Duplicate**, **Archive** and **Delete card** are at the bottom right.",
        ],
    )
    image(doc, "10-card-details", "A card with description, checklist, priority, due date and labels")
    image(doc, "11-card-attachments", "An attached file")
    image(doc, "12-card-comments", "A comment on the card")


def ai_helpers(doc):
    doc.add_heading("6. AI helpers", level=1)
    doc.add_heading("Suggest checklist steps", level=2)
    para(doc, "In a card, click **Suggest steps**. The AI reads the card and proposes steps. Untick any you don't want and click **Add steps**, or **Discard**.")
    image(doc, "13-ai-checklist", 'Suggested steps for "Write launch announcement"')
    doc.add_heading("AI assistant", level=2)
    para(
        doc,
        'Click **Ask AI** on a board. Ask anything about the board or the app, for example "What is overdue?". '
        "It can also create, move and complete cards, but only after you click **Confirm**.",
    )
    image(doc, "17-ai-assistant", "The assistant offers to move a card and waits for Confirm")
    para(
        doc,
        "The AI uses a free service with a daily limit shared by everyone. If it says it is unavailable, "
        "try again later. Your text is sent to that service only when you use an AI button.",
        muted=True,
    )


def find_cards(doc):
    doc.add_heading("7. Find cards", level=1)
    bullets(
        doc,
        [
            "**Search cards**: type part of a title or description (shortcut: **/**).",
            "**Filters**: by label, by priority, and by status (Open, Completed, Overdue, **Assigned to me**). Click **Clear filters** to see everything again.",
            "**List**: shows the cards as a table you can sort by any column. **Board** switches back.",
        ],
    )
    image(doc, "14-search", 'Searching for "page"')
    image(doc, "15-filter-priority", "Only High priority cards")
    image(doc, "16-list-view", "The list view")


def voice(doc):
    doc.add_heading("8. Voice", level=1)
    bullets(
        doc,
        [
            "**Microphone buttons** next to text boxes let you dictate instead of typing, in English or Arabic.",
            '**Voice command**: click it and say a command, for example "add card Buy domain to To Do", '
            '"move Buy domain to Done" or "show overdue". You can also type the command. '
            "Click **Confirm** to run it. **What can I say?** lists every command.",
            "The first time, your browser asks to use the microphone: click **Allow**.",
        ],
    )
    image(doc, "18-voice-command", "A voice command waiting for Confirm")


def focus_timer(doc):
    doc.add_heading("9. Focus timer", level=1)
    para(
        doc,
        "Click **Focus timer** on a board to work in focused blocks: 25 minutes of focus and a 5-minute break, "
        "or 50 and 10. Pick a card under **Focus on a card** and the time is saved on it. "
        "Use **Start focus**, **Pause**, **Reset focus**, **Reset break** and **Skip to break**. "
        "**Turn off focus timer** hides it again.",
    )
    image(doc, "19-focus-timer", "The focus timer")


def board_options(doc):
    doc.add_heading("10. Board options", level=1)
    para(doc, "The **...** button at the top right of a board opens more options:")
    bullets(
        doc,
        [
            "**Members**: share the board (see the next section).",
            "**Activity**: everything that happened on the board.",
            "**Change background**: pick a color or upload a photo (JPG, PNG or WebP, up to 10 MB).",
            "**Keyboard shortcuts**: **N** new card, **/** search, **Esc** close, **?** show the shortcuts.",
            "**Archived cards and columns**: restore or delete archived items.",
            "**Duplicate board**, **Save as template**, **Move to workspace**.",
            "**Download CSV (Excel)** or **Download JSON**: export the cards.",
            "**Auto-complete**: when on, cards moved into the last column are marked complete.",
            "**Archive board** hides it from the start page; **Delete board** removes it for good.",
        ],
    )
    image(doc, "20-board-options", "The board options menu")
    image(doc, "21-background", "Choosing a background")
    image(doc, "22-activity", "The board activity")
    image(doc, "23-archive", "Archived cards can be restored")
    image(doc, "24-keyboard-shortcuts", "Keyboard shortcuts")


def share(doc):
    doc.add_heading("11. Share and work together", level=1)
    doc.add_heading("Invite people", level=2)
    steps(
        doc,
        [
            "Click **Share** (or **...** then **Members**).",
            "Type the person's email and click **Invite**. They get an email; you can also click **Copy link** or **Share on WhatsApp** to send the invitation yourself.",
            "They open the link, sign in or create an account with that email, and click **Accept invitation**.",
        ],
    )
    image(doc, "25-share-and-members", "Invite by email, invitation link, public link and meeting link")
    image(doc, "26-invitation", "What the invited person sees")
    doc.add_heading("Permissions", level=2)
    para(doc, "You are the board **Admin**. New members start as **Viewer**. Change each member's permission in the same window:")
    simple_table(
        doc,
        ["Permission", "What they can do"],
        [
            ["Viewer", "Sees the board, changes nothing. Can ask you for more access."],
            ["Can update tasks", "Moves cards, marks them complete, ticks checklists and comments."],
            ["Full access", "Adds, edits and deletes cards and columns."],
        ],
        [4.5, 12.5],
    )
    image(doc, "27-member-permissions", "Omar can update tasks; the meeting link was sent")
    doc.add_heading("Public link", level=2)
    para(
        doc,
        "Tick **Anyone with the link can view** to get a read-only link. Anyone can open it without signing in; "
        "signed-in people can click **Join board** to become a Viewer. Untick it to stop the link.",
    )
    image(doc, "31-public-board", "A public board, view only")
    doc.add_heading("Meeting link, assigning and notifications", level=2)
    bullets(
        doc,
        [
            "Paste a Zoom, Google Meet or Teams link under **Meeting link** and click **Save link**. "
            "**Start meeting** opens it; **Send the link to members** emails it to everyone.",
            "In a card, tick a member under **Assigned to**. They get a notification and an email.",
            "The bell at the top shows your notifications; **What's new** on the start page shows them too.",
        ],
    )
    image(doc, "28-assign-card", "Assigning a card to a member")
    image(doc, "29-start-page-whats-new", "Omar's start page: the assignment and the meeting invitation")
    image(doc, "30-notifications", "The notifications window")


def account(doc):
    doc.add_heading("12. Your account", level=1)
    para(doc, "Click your initial at the top right:")
    bullets(
        doc,
        [
            "**Change password** (or **Set password** if you signed up with Google).",
            "**Email notifications: on/off**: turn reminder and notification emails on or off.",
            "**Sign out**.",
            "**Delete account**: removes your account, boards and files for good.",
        ],
    )
    image(doc, "33-account-menu", "The account menu")
    para(doc, f"The privacy policy is at **{SITE}/privacy**. Questions about your data: datasharingks@gmail.com.")


def look(doc):
    doc.add_heading("13. Theme, language and phones", level=1)
    bullets(
        doc,
        [
            "**Dark mode**: click the moon at the top right (the sun switches back).",
            "**Arabic**: click **العربية** at the top; the whole app turns right-to-left. **English** switches back. Emails and AI answers follow your language.",
            "**Phones and tablets**: everything works on small screens; columns stack and buttons wrap.",
        ],
    )
    image(doc, "34-dark-mode", "Dark mode")
    image(doc, "35-arabic", "Arabic, right to left")
    image(doc, "36-phone", "On a phone", width=Cm(6.5))


def limits(doc):
    doc.add_heading("14. Limits", level=1)
    simple_table(
        doc,
        ["What", "Limit"],
        [
            ["Cards per board", "100"],
            ["File size", "4 MB per file, 10 files per card"],
            ["Storage per account", "100 MB (guests 10 MB)"],
            ["Background photo", "10 MB (resized automatically)"],
            ["Guest accounts", "Removed after 7 days unless you sign up"],
            ["Confirmation and reset links", "Work for 24 hours"],
            ["AI", "Free service with a shared daily limit"],
        ],
        [6, 11],
    )


def help_section(doc):
    doc.add_heading("15. Help", level=1)
    simple_table(
        doc,
        ["Problem", "What to do"],
        [
            ["No confirmation email", "Look in **Spam** or **Junk**. Mark it **Not spam**. Or click **Resend email**."],
            ["The link in the email doesn't work", "It expired after 24 hours. Sign in and click **Send a new link**."],
            ["AI says it is unavailable", "The free daily limit was reached or the service is busy. Try again later."],
            ["Microphone does nothing", "Click the lock icon next to the address, allow **Microphone**, and try again."],
            ["I can only view a board", "Your permission is Viewer. Click **Request access** to ask the admin."],
        ],
        [6, 11],
    )


if __name__ == "__main__":
    print(build())
