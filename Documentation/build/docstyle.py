"""Shared look for every Word file in Documentation/: A4 pages, project colors, page numbers, helpers."""

import re
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

NAVY = RGBColor(0x03, 0x21, 0x47)
BLUE = RGBColor(0x20, 0x9D, 0xD7)
PURPLE = RGBColor(0x75, 0x39, 0x91)
GRAY = RGBColor(0x88, 0x88, 0x88)
INK = RGBColor(0x1F, 0x29, 0x37)
SHOTS = Path(__file__).parents[1] / "Screenshots"
CONTENT_WIDTH = Cm(17)
IMAGE_WIDTH = Cm(15)


def new_document(title: str, subtitle: str) -> Document:
    """A4 document with the project styles, a page-number footer and the title block."""
    doc = Document()
    section = doc.sections[0]
    section.page_width, section.page_height = Cm(21), Cm(29.7)
    section.left_margin = section.right_margin = Cm(2)
    section.top_margin = section.bottom_margin = Cm(2)
    _styles(doc)
    _page_numbers(section)
    p = doc.add_paragraph(style="Title")
    p.add_run(title)
    p = doc.add_paragraph(style="Subtitle")
    p.add_run(subtitle)
    return doc


def _styles(doc: Document) -> None:
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    normal.font.color.rgb = INK
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.15
    for name, size, color, before in (("Title", 30, NAVY, 0), ("Subtitle", 15, BLUE, 0), ("Heading 1", 18, NAVY, 18), ("Heading 2", 13, PURPLE, 12)):
        style = doc.styles[name]
        style.font.name = "Calibri"
        style.font.size = Pt(size)
        style.font.bold = name != "Subtitle"
        style.font.color.rgb = color
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(6)
        style.element.rPr.rFonts.set(qn("w:asciiTheme"), "")
    doc.styles["Heading 1"].paragraph_format.page_break_before = False


def _page_numbers(section) -> None:
    p = section.footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run()
    run.font.size = Pt(9)
    run.font.color.rgb = GRAY
    for kind, text in (("begin", None), (None, "PAGE"), ("end", None)):
        if kind:
            el = OxmlElement("w:fldChar")
            el.set(qn("w:fldCharType"), kind)
        else:
            el = OxmlElement("w:instrText")
            el.set(qn("xml:space"), "preserve")
            el.text = text
        run._r.append(el)


def add_runs(p, text: str) -> None:
    """Text with **bold** and `code` parts."""
    for part in re.split(r"(\*\*.+?\*\*|`[^`]+`)", text):
        if part.startswith("**") and part.endswith("**") and len(part) > 4:
            p.add_run(part[2:-2]).bold = True
        elif part.startswith("`") and part.endswith("`") and len(part) > 2:
            run = p.add_run(part[1:-1])
            run.font.name = "Consolas"
            run.font.size = Pt(9.5)
            run.font.color.rgb = PURPLE
        elif part:
            p.add_run(part)


def para(doc: Document, text: str, *, muted: bool = False):
    p = doc.add_paragraph()
    add_runs(p, text)
    if muted:
        for run in p.runs:
            run.font.color.rgb = GRAY
    return p


def bullets(doc: Document, items: list[str]) -> None:
    for item in items:
        add_runs(doc.add_paragraph(style="List Bullet"), item)


def steps(doc: Document, items: list[str]) -> None:
    """Numbered steps; numbers are written out so each list starts at 1."""
    for n, item in enumerate(items, 1):
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Cm(0.8)
        p.paragraph_format.first_line_indent = Cm(-0.6)
        p.paragraph_format.space_after = Pt(3)
        num = p.add_run(f"{n}.  ")
        num.bold = True
        num.font.color.rgb = PURPLE
        add_runs(p, item)


def image(doc: Document, name: str, caption: str, width=IMAGE_WIDTH) -> None:
    """A screenshot from Documentation/Screenshots with a thin border and a caption."""
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.keep_with_next = True
    pic = p.add_run().add_picture(str(SHOTS / f"{name}.png"), width=width)
    _border(pic)
    c = doc.add_paragraph()
    c.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = c.add_run(caption)
    run.italic = True
    run.font.size = Pt(9)
    run.font.color.rgb = GRAY


def _border(pic) -> None:
    sp_pr = pic._inline.graphic.graphicData.pic.spPr
    ln = OxmlElement("a:ln")
    ln.set("w", "6350")
    fill = OxmlElement("a:solidFill")
    color = OxmlElement("a:srgbClr")
    color.set("val", "D0D5DD")
    fill.append(color)
    ln.append(fill)
    sp_pr.append(ln)


def note_box(doc: Document, title: str, lines: list[str], *, border: str = "C62828", fill: str = "FDECEA") -> None:
    """A one-cell table with a colored border, used for important notes."""
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = table.cell(0, 0)
    cell.width = CONTENT_WIDTH
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = OxmlElement("w:tcBorders")
    for side in ("top", "left", "bottom", "right"):
        b = OxmlElement(f"w:{side}")
        b.set(qn("w:val"), "single")
        b.set(qn("w:sz"), "18")
        b.set(qn("w:color"), border)
        borders.append(b)
    tc_pr.append(borders)
    shade = OxmlElement("w:shd")
    shade.set(qn("w:val"), "clear")
    shade.set(qn("w:color"), "auto")
    shade.set(qn("w:fill"), fill)
    tc_pr.append(shade)
    margins = OxmlElement("w:tcMar")
    for side in ("top", "left", "bottom", "right"):
        m = OxmlElement(f"w:{side}")
        m.set(qn("w:w"), "180")
        m.set(qn("w:type"), "dxa")
        margins.append(m)
    tc_pr.append(margins)
    head = cell.paragraphs[0]
    run = head.add_run(title)
    run.bold = True
    run.font.size = Pt(13)
    run.font.color.rgb = RGBColor.from_string(border)
    for line in lines:
        add_runs(cell.add_paragraph(), line)
    doc.add_paragraph()


def simple_table(doc: Document, header: list[str], rows: list[list[str]], widths_cm: list[float]) -> None:
    """A plain grid table with a navy header row."""
    table = doc.add_table(rows=1, cols=len(header))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for cell, text, w in zip(table.rows[0].cells, header, widths_cm, strict=True):
        cell.width = Cm(w)
        run = cell.paragraphs[0].add_run(text)
        run.bold = True
        run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        shade = OxmlElement("w:shd")
        shade.set(qn("w:val"), "clear")
        shade.set(qn("w:fill"), "032147")
        cell._tc.get_or_add_tcPr().append(shade)
    for row in rows:
        cells = table.add_row().cells
        for cell, text, w in zip(cells, row, widths_cm, strict=True):
            cell.width = Cm(w)
            add_runs(cell.paragraphs[0], text)
    doc.add_paragraph()


def contents(doc: Document) -> None:
    """A table of contents field; Word fills it when the file is updated (see update_fields.ps1)."""
    doc.add_paragraph("Contents", style="Heading 2")
    run = doc.add_paragraph().add_run()
    for kind, text in (("begin", None), (None, 'TOC \\o "1-1" \\h \\z \\u'), ("separate", None), ("end", None)):
        if kind:
            el = OxmlElement("w:fldChar")
            el.set(qn("w:fldCharType"), kind)
        else:
            el = OxmlElement("w:instrText")
            el.set(qn("xml:space"), "preserve")
            el.text = text
        run._r.append(el)
