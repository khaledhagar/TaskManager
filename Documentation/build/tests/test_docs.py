"""T-620 to T-627 and T-630 for the nine Word files: run `uv run pytest` after `uv run python build_docs.py`."""

import re
from pathlib import Path

import pytest
from docx import Document

DOCS = Path(__file__).parents[2]
ROOT = DOCS.parent
PLANNING = ROOT / "Planning"
FILES = [
    "01_Overview",
    "02_Documentation",
    "03_Test_Plan",
    "04_Project_Plan",
    "05_Security",
    "06_User_Flows",
    "07_Pages",
    "08_Build_Prompts",
    "09_Phase_Checklist",
]
LONG = FILES[1:]
FLOWS = [
    "Guest start",
    "Sign up and verify",
    "Google sign-in",
    "Forgot password",
    "Create a board",
    "Cards and drag and drop",
    "Invite and accept",
    "Permissions and request access",
    "Public board and join",
    "Assign cards",
    "Meeting link",
    "AI chat actions",
    "Voice input",
    "Files",
    "Focus timer",
    "Language switch",
    "Delete account",
]


def load(name: str):
    return Document(str(DOCS / f"{name}.docx"))


def text(doc) -> str:
    cells = [c.text for t in doc.tables for r in t.rows for c in r.cells]
    return "\n".join([p.text for p in doc.paragraphs] + cells)


def headings(doc, level: int = 1) -> list[str]:
    return [p.text for p in doc.paragraphs if p.style.name == f"Heading {level}"]


@pytest.mark.parametrize("name", FILES)
def test_t620_t621_opens_with_title_headings_and_page_numbers(name):
    doc = load(name)
    assert doc.paragraphs[0].text == "Task Manager"
    assert doc.paragraphs[0].style.name == "Title"
    assert headings(doc)
    footer = doc.sections[0].footer._element.xml
    assert "PAGE" in footer
    if name in LONG:
        assert "TOC \\o" in doc.element.xml


def test_t622_overview_has_short_and_long_description():
    doc = load("01_Overview")
    paras = [p.text for p in doc.paragraphs]
    short = paras[paras.index("Short description") + 1]
    assert 1 <= len(re.findall(r"[.!?](\s|$)", short)) <= 2
    long_start = paras.index("Long description") + 1
    assert len(" ".join(paras[long_start:]).split()) >= 250


def test_t623_test_plan_lists_every_test_id_and_file():
    body = text(load("03_Test_Plan"))
    ids = set(re.findall(r"\bT-\d{3}\b", (PLANNING / "TEST_PLAN.md").read_text(encoding="utf-8")))
    assert ids and not {i for i in ids if i not in body}
    files = [*ROOT.glob("backend/tests/test_*.py"), *ROOT.glob("frontend/e2e*/*.spec.ts")]
    files += [*ROOT.glob("frontend/src/**/*.test.ts*"), *DOCS.glob("build/tests/test_*.py")]
    missing = [f.name for f in files if f.name not in body]
    assert not missing


def test_t624_project_plan_has_every_phase_and_security_every_item():
    plan = text(load("04_Project_Plan"))
    phases = re.findall(r"^## (Phase \d+ - .+)$", (PLANNING / "PLAN.md").read_text(encoding="utf-8"), re.M)
    assert phases and not [p for p in phases if p not in plan]
    security = text(load("05_Security"))
    items = set(re.findall(r"\b(?:SEC|NFR)-\d+\b", (PLANNING / "SECURITY.md").read_text(encoding="utf-8")))
    assert items and not {i for i in items if i not in security}


def test_t625_pages_have_screenshots_and_flows_are_all_covered():
    pages = load("07_Pages")
    assert len(pages.inline_shapes) >= 20
    assert len(pages.inline_shapes) >= len([h for h in headings(pages, 2) if h != "Contents"])
    flows = headings(load("06_User_Flows"), 2)
    assert not [f for f in FLOWS if not any(h.endswith(f) for h in flows)]


def test_t626_build_prompts_one_per_phase_with_goal_tasks_and_criteria():
    doc = load("08_Build_Prompts")
    phases = re.findall(r"^## (Phase \d+) - ", (PLANNING / "PLAN.md").read_text(encoding="utf-8"), re.M)
    titles = headings(doc)
    assert len([t for t in titles if t.startswith("Phase ")]) == len(phases)
    body = text(doc)
    for label in ("Goal", "Tasks", "Success criteria"):
        assert body.count(f"{label}:") >= len(phases), label


@pytest.mark.parametrize("name", [*FILES, "README"])
def test_t627_no_secrets_or_personal_addresses(name):
    body = text(load(name))
    assert not re.search(r"sk-or-v1|sk-[A-Za-z0-9]{20,}|AKIA[0-9A-Z]{12,}|eyJ[\w-]{20,}\.", body)
    assert not re.search(r"(?:password|secret|api[_ ]?key)\s*=\s*\S", body, re.I)
    emails = set(re.findall(r"[\w.+-]+@[\w-]+(?:\.[\w-]+)+", body))
    assert {e for e in emails if not e.endswith("@example.com")} <= {"datasharingks@gmail.com"}


def test_t630_phase_checklist_has_every_phase_and_task_with_its_status():
    doc = load("09_Phase_Checklist")
    assert doc.paragraphs[0].text == "Task Manager"
    plan = (PLANNING / "PLAN.md").read_text(encoding="utf-8")
    phases = re.findall(r"^## (Phase \d+ - .+)$", plan, re.M)
    assert [h for h in headings(doc) if h.startswith("Phase ")] == phases
    tasks = re.findall(r"^- \[([ xX])\] ", plan.split("\n## Phase 1 - ", 1)[1], re.M)
    marks = [p.text.split()[0] for p in doc.paragraphs if p.text.startswith(("Done ", "Open "))]
    assert marks == ["Done" if t in "xX" else "Open" for t in tasks]
