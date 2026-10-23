"""T-628 and T-627 for README.docx: run `uv run pytest` after `uv run python build_readme.py`."""

import re
from pathlib import Path

import pytest
from docx import Document
from docx.text.paragraph import Paragraph

README = Path(__file__).parents[2] / "README.docx"
SECTIONS = [
    "Getting started",
    "Your start page",
    "Create a board",
    "Work on a board",
    "Inside a card",
    "AI helpers",
    "Find cards",
    "Voice",
    "Focus timer",
    "Board options",
    "Share and work together",
    "Your account",
    "Theme, language and phones",
    "Limits",
    "Help",
]


@pytest.fixture(scope="module")
def doc():
    return Document(str(README))


def body_text(doc) -> str:
    cells = [c.text for t in doc.tables for r in t.rows for c in r.cells]
    return "\n".join([p.text for p in doc.paragraphs] + cells)


def test_first_block_under_the_title_is_the_spam_note(doc):
    texts = []
    for el in doc.element.body.iterchildren():
        if el.tag.endswith("}tbl"):
            break
        texts.append(Paragraph(el, doc).text.strip())
    assert [t for t in texts if t] == ["Task Manager", "User guide"]
    note = doc.tables[0].cell(0, 0).text
    assert "Spam" in note and "Junk" in note


def test_every_section_is_a_heading(doc):
    headings = [p.text for p in doc.paragraphs if p.style.name == "Heading 1"]
    for section in SECTIONS:
        assert any(h.endswith(section) for h in headings), section


def test_screenshots_are_included(doc):
    assert len(doc.inline_shapes) >= 30


def test_no_secrets_or_personal_addresses(doc):
    text = body_text(doc)
    assert not re.search(r"sk-or-|api[_-]?key|password\s*=|secret\s*=", text, re.I)
    emails = set(re.findall(r"[\w.+-]+@[\w-]+(?:\.[\w-]+)+", text))
    assert emails <= {"datasharingks@gmail.com", "sara@example.com", "omar@example.com"}
