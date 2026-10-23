"""08_Build_Prompts.docx: every phase of PLAN.md written as a prompt for an AI coding assistant."""

import re
from pathlib import Path

from docx.shared import Pt

from docstyle import PURPLE, add_runs, bullets, contents, new_document, para
from mdocx import plain

PLAN = Path(__file__).parents[2] / "Planning" / "PLAN.md"
TESTY = re.compile(r"\b(test|tests|tested|vitest|playwright|pytest|axe|lighthouse|verified|live check|passes|audit)\b", re.I)

GOALS = {
    1: "Set up the two projects (FastAPI backend with uv, React + TypeScript frontend with Vite) so the built site is served by FastAPI.",
    2: "Create the database tables with SQLModel and the board templates, built by one seed function.",
    3: "Add sign-up, sign-in, guest accounts and their cleanup, with a JWT in an httpOnly cookie.",
    4: "Build the REST API for boards, columns, cards, checklists and labels, with ownership checks on every call.",
    5: "Add the AI helper (LiteLLM through OpenRouter free models) and the checklist, board starter and chat endpoints.",
    6: "Add email verification, password reset and change, Google sign-in and session versioning.",
    7: "Build the frontend foundation: design tokens, API client, sign-in screens and the app shell.",
    8: "Build the board: columns and cards with drag and drop, the card window, search, filters, undo and shortcuts.",
    9: "Add file attachments on cards and board backgrounds behind one storage interface.",
    10: "Add an opt-in Pomodoro focus timer (25/5 and 50/10).",
    11: "Make the whole app work in English and Arabic (right-to-left), including server messages and emails.",
    12: "Build the AI screens: suggested steps, describe-to-AI boards and the chat panel with confirmed actions.",
    13: "Polish the interface: empty states, loading, motion, responsive layout and accessibility.",
    14: "Harden security and non-functional quality: headers, CSRF, limits, rate limits and audits.",
    15: "Write the end-to-end tests with Playwright, including accessibility and cross-browser checks.",
    16: "Finish the MVP: full test run, README, live checks and the server ready for the user.",
    17: "Add a password strength meter and refuse weak passwords on the server.",
    18: "Build board sharing in the backend: members, invitations, roles and access requests.",
    19: "Build the sharing screens: Members window, invitation page, roles and request access.",
    20: "Add public read-only boards, card assignment and a meeting link.",
    21: "Add voice input with the browser's speech recognition on the main text fields.",
    22: "Add comments with @mentions and a board activity log.",
    23: "Add notifications, reminder emails and an email opt-out.",
    24: "Add archive, duplicates, own templates, stars, WIP limits, auto-complete, covers and focus sessions.",
    25: "Add the list view and export to CSV and JSON.",
    26: "Add workspaces and the start page, then run a full regression.",
    27: "Add voice commands that act on the board after the user confirms.",
    28: "Deploy to Vercel with Supabase Postgres and Storage, and verify the live site.",
    29: "Write the documentation as Word files, the user guide, screenshots and the feature video.",
}

CONTEXT = [
    "Project: **Task Manager**, a Kanban web app: boards with columns and cards, drag and drop, guests, and a slick, professional UI.",
    "Stack: FastAPI + SQLModel (SQLite locally, Postgres in production), pwdlib + PyJWT cookie auth, LiteLLM through OpenRouter free models, "
    "React + TypeScript + Vite, dnd-kit. Tests: pytest, Vitest, Playwright.",
    "Rules: `uv run` and `uv add` for Python; small steps, each validated; write tests first; find the root cause before fixing; "
    "no over-engineering; latest library APIs; no emojis.",
    "Colors: accent yellow `#ecad0a`, blue `#209dd7`, purple `#753991` for main actions, navy `#032147` for headings, gray `#888888`.",
    "Security: ownership check on every endpoint (404 for other people's data), httpOnly SameSite cookie, Origin check, input limits, "
    "rate limits on auth, guests and AI, strict CSP, AI actions only after user confirmation, API key never sent to the browser.",
]


def phases() -> list[tuple[int, str, list[str]]]:
    text = PLAN.read_text(encoding="utf-8")
    found = []
    for match in re.finditer(r"^## Phase (\d+) - (.+?)\n(.*?)(?=^## |\Z)", text, re.M | re.S):
        items = [plain(i) for i in re.findall(r"^- \[[ xX]\] (.+)$", match.group(3), re.M)]
        found.append((int(match.group(1)), match.group(2).strip(), items))
    return found


def label(doc, name: str) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(6)
    run = p.add_run(f"{name}:")
    run.bold = True
    run.font.color.rgb = PURPLE


def build(out: Path) -> None:
    doc = new_document("Task Manager", "Build prompts")
    para(
        doc,
        "Each phase below is a ready-to-use prompt for an AI coding assistant. Give it the shared context first, "
        "then one phase at a time, and check the success criteria before moving on.",
    )
    contents(doc)
    doc.add_heading("Shared context (give this first)", level=1)
    bullets(doc, CONTEXT)
    for number, title, items in phases():
        doc.add_heading(f"Phase {number} - {title}", level=1)
        label(doc, "Goal")
        add_runs(doc.add_paragraph(), GOALS[number])
        tasks = [i for i in items if not TESTY.search(i)]
        checks = [i for i in items if TESTY.search(i)]
        label(doc, "Tasks")
        bullets(doc, tasks or ["See the success criteria."])
        label(doc, "Success criteria")
        bullets(doc, [*checks, "Every task above is done; `uv run pytest`, `npx vitest run`, lint and type checks pass."])
    doc.save(str(out))
